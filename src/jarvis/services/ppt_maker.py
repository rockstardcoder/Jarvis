import os
import re
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


PRESENTATION_DIR = Path(
    "data/presentations"
)


THEMES = {
    "professional": {
        "background": RGBColor(10, 22, 40),
        "panel": RGBColor(18, 38, 64),
        "panel_light": RGBColor(235, 242, 250),
        "text": RGBColor(245, 248, 252),
        "muted": RGBColor(180, 196, 215),
        "dark_text": RGBColor(25, 35, 50),
        "accent": RGBColor(0, 170, 255),
        "accent_2": RGBColor(50, 220, 180),
        "white": RGBColor(255, 255, 255),
    },
    "school": {
        "background": RGBColor(245, 248, 255),
        "panel": RGBColor(225, 235, 255),
        "panel_light": RGBColor(255, 255, 255),
        "text": RGBColor(20, 40, 70),
        "muted": RGBColor(80, 100, 130),
        "dark_text": RGBColor(20, 40, 70),
        "accent": RGBColor(30, 90, 200),
        "accent_2": RGBColor(0, 150, 120),
        "white": RGBColor(255, 255, 255),
    },
    "simple": {
        "background": RGBColor(255, 255, 255),
        "panel": RGBColor(242, 245, 250),
        "panel_light": RGBColor(250, 250, 250),
        "text": RGBColor(30, 30, 30),
        "muted": RGBColor(90, 90, 90),
        "dark_text": RGBColor(30, 30, 30),
        "accent": RGBColor(40, 100, 220),
        "accent_2": RGBColor(80, 160, 120),
        "white": RGBColor(255, 255, 255),
    }
}


class PPTMaker:

    def ensure_output_folder(
        self
    ):
        PRESENTATION_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

    def sanitize_filename(
        self,
        text: str
    ) -> str:
        text = str(text).strip().lower()
        text = re.sub(
            r"[^a-z0-9]+",
            "_",
            text
        )
        text = text.strip("_")

        if not text:
            text = "presentation"

        return text[:60]

    def clean_topic(
        self,
        topic: str
    ) -> str:
        topic = str(topic).strip()

        if not topic:
            return "Untitled Presentation"

        return topic[0].upper() + topic[1:]

    def create_output_path(
        self,
        topic: str
    ) -> Path:
        self.ensure_output_folder()

        timestamp = datetime.now().strftime(
            "%Y-%m-%d_%H-%M-%S"
        )

        filename = (
            f"{self.sanitize_filename(topic)}_"
            f"{timestamp}.pptx"
        )

        return PRESENTATION_DIR / filename

    def set_slide_background(
        self,
        slide,
        color
    ):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_textbox(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        width: float,
        height: float,
        size: int = 24,
        color=None,
        bold: bool = False,
        align=None,
        font_name: str = "Aptos",
    ):
        box = slide.shapes.add_textbox(
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height)
        )

        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.08)
        frame.margin_right = Inches(0.08)
        frame.margin_top = Inches(0.03)
        frame.margin_bottom = Inches(0.03)

        paragraph = frame.paragraphs[0]
        paragraph.text = text

        if align is not None:
            paragraph.alignment = align

        run = paragraph.runs[0]
        run.font.name = font_name
        run.font.size = Pt(size)
        run.font.bold = bold

        if color is not None:
            run.font.color.rgb = color

        return box

    def add_bullets(
        self,
        slide,
        bullets: list[str],
        x: float,
        y: float,
        width: float,
        height: float,
        size: int,
        color,
        line_spacing: float = 0.95
    ):
        box = slide.shapes.add_textbox(
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height)
        )

        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.1)
        frame.margin_right = Inches(0.1)

        for index, bullet in enumerate(bullets):
            if index == 0:
                paragraph = frame.paragraphs[0]
            else:
                paragraph = frame.add_paragraph()

            paragraph.text = f"• {bullet}"
            paragraph.space_after = Pt(8)
            paragraph.line_spacing = line_spacing

            run = paragraph.runs[0]
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.color.rgb = color

        return box

    def add_footer(
        self,
        slide,
        slide_number: int,
        topic: str,
        theme: dict
    ):
        self.add_textbox(
            slide,
            topic,
            0.55,
            7.05,
            7.5,
            0.25,
            size=9,
            color=theme["muted"]
        )

        self.add_textbox(
            slide,
            str(slide_number),
            12.25,
            7.05,
            0.5,
            0.25,
            size=9,
            color=theme["muted"],
            align=PP_ALIGN.RIGHT
        )

    def add_header(
        self,
        slide,
        title: str,
        slide_number: int,
        topic: str,
        theme: dict
    ):
        self.add_textbox(
            slide,
            title,
            0.65,
            0.35,
            9.8,
            0.45,
            size=24,
            color=theme["text"],
            bold=True,
            font_name="Aptos Display"
        )

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.65),
            Inches(0.95),
            Inches(1.15),
            Inches(0.06)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = theme["accent"]
        line.line.fill.background()

        self.add_footer(
            slide,
            slide_number,
            topic,
            theme
        )

    def add_card(
        self,
        slide,
        x: float,
        y: float,
        width: float,
        height: float,
        title: str,
        body: str,
        theme: dict,
        number: str | None = None
    ):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height)
        )

        card.fill.solid()
        card.fill.fore_color.rgb = theme["panel"]
        card.line.color.rgb = theme["accent"]

        if number:
            self.add_textbox(
                slide,
                number,
                x + 0.2,
                y + 0.15,
                0.45,
                0.35,
                size=14,
                color=theme["accent"],
                bold=True,
                align=PP_ALIGN.CENTER
            )

            title_x = x + 0.75
            title_width = width - 0.95

        else:
            title_x = x + 0.25
            title_width = width - 0.5

        self.add_textbox(
            slide,
            title,
            title_x,
            y + 0.15,
            title_width,
            0.35,
            size=15,
            color=theme["text"],
            bold=True
        )

        self.add_textbox(
            slide,
            body,
            x + 0.25,
            y + 0.65,
            width - 0.5,
            height - 0.85,
            size=11,
            color=theme["muted"]
        )

    def build_content(
        self,
        topic: str
    ) -> dict:
        topic = self.clean_topic(
            topic
        )

        return {
            "title": topic,
            "subtitle": (
                "A structured, professional overview prepared by Jarvis"
            ),
            "agenda": [
                "Understand the core idea and context",
                "Explore the key points that matter most",
                "Review practical applications and strategy",
                "Identify benefits, risks, and next steps"
            ],
            "overview": [
                f"{topic} is an important area that can be understood through its purpose, impact, and practical use.",
                "A clear overview helps connect the topic to real-world decisions and outcomes.",
                "The best approach is to break the subject into core ideas, examples, challenges, and solutions."
            ],
            "key_points": [
                (
                    "Purpose",
                    f"Clarifies why {topic} matters and what problem or need it addresses."
                ),
                (
                    "Process",
                    "Explains how the concept works, develops, or is applied in real situations."
                ),
                (
                    "Impact",
                    "Shows the benefits, limitations, and possible long-term outcomes."
                )
            ],
            "strategy": [
                "Define the objective clearly",
                "Collect relevant information and examples",
                "Organize ideas into logical sections",
                "Evaluate benefits, risks, and improvements",
                "Present conclusions with practical next steps"
            ],
            "benefits": [
                "Improves understanding and decision-making",
                "Creates a structured way to compare ideas",
                "Helps communicate complex information clearly",
                "Supports planning, learning, and presentation quality"
            ],
            "challenges": [
                (
                    "Information overload",
                    "Use clear sections and focus only on the most relevant points."
                ),
                (
                    "Lack of examples",
                    "Add realistic examples, use cases, or comparisons."
                ),
                (
                    "Weak conclusion",
                    "End with clear takeaways and action points."
                )
            ],
            "conclusion": [
                f"{topic} becomes easier to understand when it is explained through structure and examples.",
                "The most effective presentations focus on clarity, relevance, and practical value.",
                "Next step: add specific data, images, charts, or case studies for a deeper version."
            ]
        }

    def create_title_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(0.22),
            Inches(7.5)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = theme["accent"]
        accent.line.fill.background()

        self.add_textbox(
            slide,
            "JARVIS PRESENTATION",
            0.75,
            0.65,
            4.0,
            0.35,
            size=12,
            color=theme["accent"],
            bold=True
        )

        self.add_textbox(
            slide,
            content["title"],
            0.75,
            2.05,
            8.8,
            1.35,
            size=40,
            color=theme["text"],
            bold=True,
            font_name="Aptos Display"
        )

        self.add_textbox(
            slide,
            content["subtitle"],
            0.8,
            3.55,
            7.2,
            0.65,
            size=17,
            color=theme["muted"]
        )

        date_text = datetime.now().strftime(
            "%d %B %Y"
        )

        self.add_textbox(
            slide,
            date_text,
            0.8,
            6.75,
            3.0,
            0.3,
            size=11,
            color=theme["muted"]
        )

        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(10.2),
            Inches(1.3),
            Inches(2.2),
            Inches(2.2)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = theme["accent"]
        circle.line.fill.background()

        self.add_textbox(
            slide,
            "01",
            10.7,
            1.95,
            1.2,
            0.55,
            size=28,
            color=theme["white"],
            bold=True,
            align=PP_ALIGN.CENTER
        )

    def create_agenda_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        self.add_header(
            slide,
            "Agenda",
            2,
            content["title"],
            theme
        )

        for index, item in enumerate(content["agenda"]):
            y = 1.55 + index * 1.1

            self.add_card(
                slide,
                1.0,
                y,
                10.9,
                0.75,
                f"Step {index + 1}",
                item,
                theme,
                number=str(index + 1)
            )

    def create_overview_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        self.add_header(
            slide,
            "Overview",
            3,
            content["title"],
            theme
        )

        self.add_bullets(
            slide,
            content["overview"],
            0.95,
            1.55,
            6.1,
            3.8,
            size=18,
            color=theme["text"]
        )

        self.add_card(
            slide,
            7.55,
            1.55,
            4.3,
            3.8,
            "Main Idea",
            (
                "This slide introduces the subject, explains why it matters, "
                "and sets up the rest of the presentation."
            ),
            theme
        )

    def create_key_points_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        self.add_header(
            slide,
            "Key Points",
            4,
            content["title"],
            theme
        )

        for index, item in enumerate(content["key_points"]):
            title, body = item

            self.add_card(
                slide,
                0.75 + index * 4.1,
                1.65,
                3.65,
                3.75,
                title,
                body,
                theme,
                number=str(index + 1)
            )

    def create_strategy_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        self.add_header(
            slide,
            "Practical Approach",
            5,
            content["title"],
            theme
        )

        for index, step in enumerate(content["strategy"]):
            x = 0.85 + index * 2.45

            bubble = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(x),
                Inches(1.75),
                Inches(0.72),
                Inches(0.72)
            )
            bubble.fill.solid()
            bubble.fill.fore_color.rgb = theme["accent"]
            bubble.line.fill.background()

            self.add_textbox(
                slide,
                str(index + 1),
                x,
                1.92,
                0.72,
                0.3,
                size=13,
                color=theme["white"],
                bold=True,
                align=PP_ALIGN.CENTER
            )

            self.add_textbox(
                slide,
                step,
                x - 0.15,
                2.75,
                1.8,
                1.15,
                size=12,
                color=theme["text"],
                bold=True,
                align=PP_ALIGN.CENTER
            )

        self.add_textbox(
            slide,
            "A strong presentation follows a clear path: define, explain, evaluate, and conclude.",
            1.1,
            5.35,
            10.5,
            0.65,
            size=18,
            color=theme["muted"],
            align=PP_ALIGN.CENTER
        )

    def create_benefits_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        self.add_header(
            slide,
            "Benefits and Impact",
            6,
            content["title"],
            theme
        )

        self.add_bullets(
            slide,
            content["benefits"],
            1.0,
            1.45,
            10.6,
            4.5,
            size=20,
            color=theme["text"]
        )

    def create_challenges_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        self.add_header(
            slide,
            "Challenges and Solutions",
            7,
            content["title"],
            theme
        )

        for index, item in enumerate(content["challenges"]):
            title, body = item

            self.add_card(
                slide,
                1.0,
                1.45 + index * 1.35,
                10.8,
                0.95,
                title,
                body,
                theme,
                number=str(index + 1)
            )

    def create_conclusion_slide(
        self,
        prs,
        content,
        theme
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.set_slide_background(
            slide,
            theme["background"]
        )

        self.add_header(
            slide,
            "Conclusion",
            8,
            content["title"],
            theme
        )

        self.add_bullets(
            slide,
            content["conclusion"],
            1.0,
            1.6,
            10.7,
            3.8,
            size=20,
            color=theme["text"]
        )

        self.add_textbox(
            slide,
            "Thank You",
            1.0,
            5.85,
            10.7,
            0.55,
            size=26,
            color=theme["accent"],
            bold=True,
            align=PP_ALIGN.CENTER,
            font_name="Aptos Display"
        )

    def create_presentation(
        self,
        topic: str,
        style: str = "professional"
    ) -> str:
        topic = self.clean_topic(
            topic
        )

        style = str(style).lower().strip()

        if style not in THEMES:
            style = "professional"

        theme = THEMES[style]
        content = self.build_content(
            topic
        )

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        self.create_title_slide(
            prs,
            content,
            theme
        )

        self.create_agenda_slide(
            prs,
            content,
            theme
        )

        self.create_overview_slide(
            prs,
            content,
            theme
        )

        self.create_key_points_slide(
            prs,
            content,
            theme
        )

        self.create_strategy_slide(
            prs,
            content,
            theme
        )

        self.create_benefits_slide(
            prs,
            content,
            theme
        )

        self.create_challenges_slide(
            prs,
            content,
            theme
        )

        self.create_conclusion_slide(
            prs,
            content,
            theme
        )

        output_path = self.create_output_path(
            topic
        )

        prs.save(
            output_path
        )

        return f"Presentation created: {output_path}"

    def open_presentation_folder(
        self
    ) -> str:
        try:
            self.ensure_output_folder()

            os.startfile(
                str(PRESENTATION_DIR.resolve())
            )

            return "Opening presentations folder."

        except Exception as e:
            return f"Failed to open presentations folder: {e}"